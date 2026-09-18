"""Render the editable participant deck and a matching vector PDF.

Install the optional tooling in presentations/requirements.txt, then run this file.
The JSON content is the editable source; no test-environment file is read.
"""

from pathlib import Path
import json
import math
import os

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor

ROOT = Path(__file__).resolve().parent
W, H = 960, 540
C = dict(
    bg="F5F8F7",
    paper="FFFFFF",
    ink="173A42",
    muted="557079",
    teal="008C7B",
    lime="CAEFA1",
    line="DCE6E2",
    pale="E8F3ED",
    dark="102F37",
    white="FFFFFF",
    blue="5E8DCD",
)
FONTS = Path(
    os.environ.get("VENDING_SLIDE_FONT_DIR", "/usr/share/fonts/truetype/dejavu")
)
for key, file in [
    ("regular", "DejaVuSans.ttf"),
    ("bold", "DejaVuSans-Bold.ttf"),
    ("mono", "DejaVuSansMono.ttf"),
]:
    pdfmetrics.registerFont(TTFont(key, str(FONTS / file)))


def wrap(text, font, size, width):
    lines = []
    for paragraph in text.split("\n"):
        if not paragraph:
            lines.append("")
            continue
        if pdfmetrics.stringWidth(paragraph, font, size) <= width:
            lines.append(paragraph)
            continue
        line = ""
        for word in paragraph.split(" "):
            candidate = f"{line} {word}" if line else word
            if pdfmetrics.stringWidth(candidate, font, size) <= width:
                line = candidate
            else:
                if line:
                    lines.append(line)
                if pdfmetrics.stringWidth(word, font, size) > width:
                    raise ValueError(f"Unbreakable text exceeds its box: {word}")
                line = word
        lines.append(line)
    return lines


class Deck:
    def __init__(self, output):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Pt(W), Pt(H)
        self.prs.core_properties.title = "Vending Bench — Harness Design Challenge"
        self.prs.core_properties.subject = (
            "Participant briefing for a pre-generated test environment"
        )
        self.prs.core_properties.author = "Vending Bench"
        self.pdf = canvas.Canvas(str(output.with_suffix(".pdf")), pagesize=(W, H))
        self.pdf.setTitle(self.prs.core_properties.title)
        self.pdf.setAuthor("Vending Bench")
        self.output = output
        self.page = 0
        self.text_boxes = []

    def color(self, name):
        return C.get(name, name)

    def rect(self, x, y, w, h, fill="paper", radius=0, stroke=None):
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
            Pt(x),
            Pt(y),
            Pt(w),
            Pt(h),
        )
        if radius:
            shape.adjustments[0] = min(0.15, radius / min(w, h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(self.color(fill))
        if stroke:
            shape.line.color.rgb = RGBColor.from_string(self.color(stroke))
            shape.line.width = Pt(0.7)
        else:
            shape.line.fill.background()
        self.pdf.setFillColor(HexColor("#" + self.color(fill)))
        self.pdf.setStrokeColor(HexColor("#" + self.color(stroke or fill)))
        self.pdf.setLineWidth(0.7)
        self.pdf.roundRect(
            x, H - y - h, w, h, radius, fill=1, stroke=1 if stroke else 0
        )
        return shape

    def line(self, x1, y1, x2, y2, color="line", width=1):
        s = self.slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2)
        )
        s.line.color.rgb = RGBColor.from_string(self.color(color))
        s.line.width = Pt(width)
        self.pdf.setStrokeColor(HexColor("#" + self.color(color)))
        self.pdf.setLineWidth(width)
        self.pdf.line(x1, H - y1, x2, H - y2)

    def text(
        self,
        text,
        x,
        y,
        w,
        h,
        size=18,
        color="ink",
        bold=False,
        mono=False,
        leading=1.27,
    ):
        font = "mono" if mono else "bold" if bold else "regular"
        lines = wrap(text, font, size, w)
        used = size + (len(lines) - 1) * size * leading
        if used > h + 0.1:
            raise ValueError(
                f"Slide {self.page}: text overflow ({used:.1f}>{h}): {text}"
            )
        if min(x, y) < 0 or x + w > W + 0.1 or y + h > H + 0.1:
            raise ValueError(f"Slide {self.page}: box outside slide")
        box = self.slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        tf = box.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = "DejaVu Sans Mono" if mono else "DejaVu Sans"
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = RGBColor.from_string(self.color(color))
            p.space_before = Pt(0)
            p.space_after = Pt(0)
            p.line_spacing = Pt(size * leading)
        self.pdf.setFont(font, size)
        self.pdf.setFillColor(HexColor("#" + self.color(color)))
        for i, line in enumerate(lines):
            self.pdf.drawString(x, H - y - size * 0.82 - i * size * leading, line)
        self.text_boxes.append(
            dict(
                slide=self.page,
                text=text,
                x=x,
                y=y,
                w=w,
                h=h,
                lines=len(lines),
                size=size,
            )
        )
        return used

    def start(self, spec, index):
        self.page = index
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.dark = spec["layout"] in ("cover", "score", "checklist")
        self.rect(0, 0, W, H, "dark" if self.dark else "bg")
        self.text(
            spec["section"], 52, 28, 810, 18, 10, "lime" if self.dark else "teal", True
        )
        if spec["layout"] != "cover":
            self.text(
                spec["title"],
                52,
                65,
                856,
                90,
                31,
                "white" if self.dark else "ink",
                True,
                leading=1.16,
            )
        self.slide.notes_slide.notes_text_frame.text = (
            spec["notes"]
            + "\n\nImplementation references: "
            + ", ".join(spec["sources"])
        )

    def finish(self, spec):
        if spec["layout"] != "cover":
            self.rect(52, 460, 856, 43, "193F45" if self.dark else "pale", radius=7)
            self.text(
                spec["takeaway"],
                66,
                472,
                827,
                28,
                12,
                "lime" if self.dark else "ink",
                bold=True,
                leading=1.15,
            )
        self.line(52, 517, 908, 517, "32515A" if self.dark else "line", 0.65)
        self.text(
            "VENDING BENCH  /  PARTICIPANT BRIEFING",
            52,
            524,
            670,
            10,
            7.5,
            "A6C1C6" if self.dark else "muted",
        )
        self.text(
            f"{self.page:02}",
            881,
            523,
            27,
            12,
            9,
            "lime" if self.dark else "teal",
            True,
        )
        self.pdf.showPage()

    def bullet(self, text, x, y, w, size=17, color="ink"):
        self.rect(x, y + 6, 5, 5, "lime" if self.dark else "teal", radius=2)
        return self.text(text, x + 18, y, w - 18, 110, size, color, leading=1.35)

    def columns(self, spec, top=174, height=264):
        for i, col in enumerate(spec["columns"]):
            x = 52 + i * 438
            self.rect(x, top, 418, height, "paper", radius=10)
            self.rect(x, top, 418, 4, "teal" if i == 0 else "blue")
            self.text(
                col["title"],
                x + 22,
                top + 22,
                374,
                35,
                12,
                "teal" if i == 0 else "muted",
                True,
            )
            y = top + 62
            for item in col["items"]:
                y += self.bullet(item, x + 22, y, 373, 16) + 12
            if y - 12 > top + height - 10:
                raise ValueError(f"Column content too tall on slide {self.page}")

    def steps(self, spec, loop=False):
        steps = spec["steps"]
        gap = 17
        w = (856 - gap * (len(steps) - 1)) / len(steps)
        top = 185
        for i, step in enumerate(steps):
            x = 52 + i * (w + gap)
            self.rect(x, top, w, 234 if not loop else 202, "paper", radius=10)
            self.rect(x + 18, top + 20, 34, 30, "pale", radius=7)
            self.text(f"{i+1:02}", x + 25, top + 27, 26, 18, 12, "teal", True)
            self.text(
                step["title"],
                x + 18,
                top + 76,
                w - 36,
                33,
                22 if not loop else 18,
                "ink",
                True,
            )
            self.text(
                step["body"],
                x + 18,
                top + 120,
                w - 36,
                105 if not loop else 73,
                17 if not loop else 15,
                leading=1.33,
            )
            if i < len(steps) - 1:
                self.text("›", x + w + 4, top + 98, 14, 30, 25, "teal", True)
        if loop:
            self.line(76, 417, 260, 417, "teal", 2)
            self.line(705, 417, 884, 417, "teal", 2)
            self.text(
                "Repeat while running and within budget",
                275,
                402,
                420,
                32,
                15,
                "teal",
                True,
            )

    def stats(self, spec, budget=False):
        items = spec["stats"]
        gap = 16
        w = (856 - gap * (len(items) - 1)) / len(items)
        for i, item in enumerate(items):
            x = 52 + i * (w + gap)
            self.rect(x, 177, w, 125, "paper", radius=10)
            self.text(
                item["value"],
                x + 17,
                193,
                w - 34,
                49,
                32 if budget else 27,
                "teal",
                True,
            )
            self.text(item["label"], x + 17, 251, w - 34, 43, 13.5, "muted")
        y = 325
        for item in spec["bullets"]:
            y += self.bullet(item, 58, y, 842, 17) + 18

    def table(self, spec):
        count = len(spec["headers"])
        widths = (
            [220, 636]
            if count == 2
            else ([165, 421, 270] if self.page == 11 else [195, 320, 341])
        )
        x = 52
        top = 167
        self.rect(x, top, 856, 34, "ink", radius=5)
        left = x
        for header, width in zip(spec["headers"], widths):
            self.text(
                header.upper(), left + 13, top + 10, width - 25, 18, 12, "white", True
            )
            left += width
        rows = spec["rows"]
        height = 242 / len(rows)
        for i, row in enumerate(rows):
            y = top + 34 + i * height
            self.rect(x, y, 856, height, "paper" if i % 2 == 0 else "EDF3F0")
            left = x
            for j, (text, width) in enumerate(zip(row, widths)):
                size = 15.5
                self.text(
                    text,
                    left + 13,
                    y + 12,
                    width - 25,
                    height - 14,
                    size,
                    "teal" if j == 0 else "ink",
                    bold=j == 0,
                    leading=1.19,
                )
                left += width

    def graph(self, spec):
        diagram_x = 52
        diagram_y = 170
        diagram_w = 612
        diagram_h = 266
        side_x = 684
        side_w = 224
        self.rect(diagram_x, diagram_y, diagram_w, diagram_h, "paper", radius=10)
        self.rect(side_x, diagram_y, 224, diagram_h, "paper", radius=10)
        self.text(
            spec.get("graph_title", "Harness graph"),
            diagram_x + 20,
            diagram_y + 18,
            diagram_w - 40,
            22,
            10.5,
            "muted",
            True,
        )
        self.text(
            spec.get("sidebar_title", "Behavior rules"),
            side_x + 18,
            diagram_y + 18,
            190,
            22,
            10.5,
            "teal",
            True,
        )

        nodes = {node["id"]: node for node in spec["nodes"]}
        for edge in spec["edges"]:
            if isinstance(edge, dict):
                src = edge["from"]
                dst = edge["to"]
                color = edge.get("color", "line")
            else:
                src, dst = edge
                color = "line"
            a = nodes[src]
            b = nodes[dst]
            ax = diagram_x + a["x"]
            ay = diagram_y + a["y"]
            bx = diagram_x + b["x"]
            by = diagram_y + b["y"]
            aw = a["w"]
            ah = a["h"]
            bw = b["w"]
            bh = b["h"]
            if bx >= ax + aw:
                x1 = ax + aw
                y1 = ay + ah / 2
                x2 = bx
                y2 = by + bh / 2
            elif ax >= bx + bw:
                x1 = ax
                y1 = ay + ah / 2
                x2 = bx + bw
                y2 = by + bh / 2
            elif by >= ay + ah:
                x1 = ax + aw / 2
                y1 = ay + ah
                x2 = bx + bw / 2
                y2 = by
            else:
                x1 = ax + aw / 2
                y1 = ay
                x2 = bx + bw / 2
                y2 = by + bh
            self.line(x1, y1, x2, y2, color, 1.8 if color != "line" else 1.2)

        for node in spec["nodes"]:
            x = diagram_x + node["x"]
            y = diagram_y + node["y"]
            fill = node.get("fill", "paper")
            stroke = node.get("stroke", "line")
            radius = node.get("radius", 8)
            self.rect(x, y, node["w"], node["h"], fill, radius=radius, stroke=stroke)
            self.text(
                node["label"],
                x + 10,
                y + 10,
                node["w"] - 20,
                node["h"] - 18,
                node.get("size", 14),
                node.get("color", "ink"),
                bold=node.get("bold", True),
                leading=node.get("leading", 1.08),
            )

        y = diagram_y + 52
        for item in spec.get("bullets", []):
            y += self.bullet(item, side_x + 14, y, 194, 13.5) + 11

    def render(self, spec, index):
        self.start(spec, index)
        layout = spec["layout"]
        if layout == "cover":
            self.rect(52, 76, 60, 5, "lime")
            self.text(spec["title"], 52, 117, 610, 172, 46, "white", True, leading=1.19)
            self.text(spec["subtitle"], 55, 301, 530, 43, 19, "lime")
            self.text(spec["takeaway"], 55, 363, 510, 77, 21, "white", leading=1.4)
            self.rect(674, 88, 218, 345, "193F45", radius=18, stroke="35535A")
            self.rect(693, 113, 180, 33, "lime", radius=7)
            self.text("YOUR AGENT", 709, 123, 160, 17, 12, "ink", True)
            for row in range(4):
                for col in range(3):
                    x = 694 + col * 60
                    y = 168 + row * 47
                    self.rect(x, y, 48, 36, "294E54", radius=5)
                    self.rect(
                        x + 17,
                        y + 8,
                        15,
                        22,
                        "lime" if (row + col) % 3 == 0 else "71A99B",
                        radius=4,
                    )
            self.rect(696, 374, 120, 16, "dark", radius=4)
            self.rect(838, 371, 20, 22, "teal", radius=4)
        elif layout == "columns":
            self.columns(spec)
        elif layout == "steps":
            self.steps(spec)
        elif layout == "loop":
            self.steps(spec, True)
        elif layout == "stats":
            self.stats(spec)
        elif layout == "budget":
            self.stats(spec, True)
        elif layout == "table":
            self.table(spec)
        elif layout == "score":
            self.text("FINAL SCORE", 55, 169, 800, 19, 11, "lime", True)
            self.text(
                spec["formula"], 55, 205, 845, 79, 25, "white", True, leading=1.35
            )
            for i, label in enumerate(spec["example"]):
                x = 52 + i * 174
                self.rect(x, 326, 160, 98, "lime" if i == 4 else "22464D", radius=9)
                parts = label.split(" ")
                amount = (
                    " ".join(parts[:2]) if parts[0] in ("+", "−", "=") else parts[0]
                )
                desc = (
                    " ".join(parts[2:])
                    if parts[0] in ("+", "−", "=")
                    else " ".join(parts[1:])
                )
                self.text(
                    amount, x + 15, 342, 136, 40, 27, "ink" if i == 4 else "white", True
                )
                self.text(desc, x + 15, 391, 136, 20, 13, "ink" if i == 4 else "A6C1C6")
        elif layout == "pricing":
            self.rect(52, 170, 373, 262, "paper", radius=10)
            self.text("QUALITATIVE RELATIONSHIP", 72, 186, 329, 22, 10, "muted", True)
            self.line(105, 385, 393, 385, "muted", 1)
            self.line(105, 385, 105, 227, "muted", 1)
            self.text("Expected demand", 73, 213, 275, 23, 12, "teal")
            self.text("Selling price →", 201, 400, 194, 22, 12, "muted")
            pts = [
                (114 + i * 10, 245 + 124 * (1 - math.exp(-i / 8))) for i in range(27)
            ]
            for a, b in zip(pts, pts[1:]):
                self.line(*a, *b, "teal", 3)
            y = 173
            for item in spec["bullets"]:
                y += self.bullet(item, 458, y, 450, 17) + 18
        elif layout == "code":
            self.rect(52, 174, 575, 262, "ink", radius=10)
            self.text("LOCAL PRACTICE", 70, 190, 532, 18, 10, "lime", True)
            self.text(
                spec["code"], 70, 225, 537, 204, 14, "white", mono=True, leading=1.53
            )
            y = 184
            for item in spec["bullets"]:
                y += self.bullet(item, 653, y, 255, 15.5) + 18
        elif layout == "review":
            self.rect(52, 164, 856, 42, "ink", radius=8)
            self.text("RUN EXPLORER", 71, 178, 260, 18, 11, "lime", True)
            self.text(
                "http://localhost:8080", 659, 177, 225, 23, 12, "white", mono=True
            )
            self.columns(spec, top=222, height=215)
        elif layout == "graph":
            self.graph(spec)
        elif layout == "checklist":
            for i, item in enumerate(spec["items"]):
                y = 179 + i * 49
                self.rect(54, y, 27, 27, "lime", radius=5)
                self.text(f"{i+1}", 63, y + 5, 18, 21, 13, "ink", True)
                self.text(item, 102, y + 2, 795, 39, 20, "white", leading=1.2)
        elif layout == "example":
            for i, (label, key) in enumerate(
                [("REQUEST", "request"), ("RESPONSE EXCERPT", "response")]
            ):
                x = 52 + i * 438
                self.rect(x, 164, 418, 279, "ink", radius=10)
                self.text(label, x + 18, 182, 380, 19, 10, "lime", True)
                self.text(
                    spec[key],
                    x + 18,
                    215,
                    382,
                    220,
                    13.5,
                    "white",
                    mono=True,
                    leading=1.37,
                )
        else:
            raise ValueError(layout)
        self.finish(spec)

    def save(self):
        self.prs.save(self.output.with_suffix(".pptx"))
        self.pdf.save()
        (ROOT / "layout-check.json").write_text(
            json.dumps(
                dict(
                    slides=self.page,
                    text_boxes=len(self.text_boxes),
                    bounds_checked=True,
                ),
                indent=2,
            )
            + "\n"
        )


def main():
    spec = json.loads((ROOT / "participant-briefing.json").read_text())
    notes = [
        "# Participant briefing — speaker notes",
        "",
        f"{len(spec['slides'])} slides total. Approximate speaking time: 20 minutes plus questions.",
        "",
    ]
    for i, slide in enumerate(spec["slides"], 1):
        notes += [
            f"## {i:02}. {slide['title'].replace(chr(10),' ')}",
            "",
            slide["notes"],
            "",
            "Sources: " + ", ".join(slide["sources"]),
            "",
        ]
    (ROOT / "participant-speaker-notes.md").write_text("\n".join(notes))
    deck = Deck(ROOT / "participant-briefing")
    for i, slide in enumerate(spec["slides"], 1):
        deck.render(slide, i)
    deck.save()
    print(f"Created {len(spec['slides'])} slides as editable PPTX and PDF")


if __name__ == "__main__":
    main()
